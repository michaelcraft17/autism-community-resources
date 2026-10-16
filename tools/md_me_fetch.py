#!/usr/bin/env python3
"""
Moldova and Montenegro -- registers of non-commercial organisations (open data, CKAN, keyless).

Moldova: Public Services Agency "Registrul de stat al unitatilor de drept - organizatii necomerciale" (weekly XLSX,
         dataset.gov.md; newest resource taken from package_show). Active rows only (no liquidation date).
Montenegro: list of active NGOs (data.gov.me, newest resource of the dataset). The authorised-person columns are never read.

Selection: names matching disability terms in Romanian / Montenegrin. Moldova's "Conducator" (a named person) is never read.
Addresses with an apartment / office number ("ap.(of.)") are shown at locality level; others are geocoded at street level
(Nominatim, cached) with a locality fallback.
Output: new_resources_moldova.json, new_resources_montenegro.json
"""
import json
import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import registry_common as rc  # noqa: E402

SCRATCH = os.path.join(rc.REPO, "tools", "small_scratch")
UA = {"User-Agent": rc.UA}
MD_PKG = "https://dataset.gov.md/api/3/action/package_show?id=date-din-registrul-de-stat-al-unitatilor-de-drept-privind-organizatiile-necomerciale"
ME_PKG = "https://data.gov.me/api/3/action/package_show?id=spisak-aktivnih-nevladinih-udruzenja-januar-2025"
MD_RE = re.compile(r"autis|asperger|dizabilit|handicap|nev[aă]z[aă]tori|\borbi\b|\bsurzi|hipoacuz|epilepsi|cerin[tț]e speciale|nevoi speciale|"
                   r"olimpiada special|sindrom\w* down|paralizie cerebral|retard mental|deficien[tț]", re.I)
ME_SKIP = re.compile(r"invalid\w* rada|invalid\w* rata|vojn\w+ invalid|ratn\w+ vojn|boraca|radne sposobnosti|rada\b", re.I)
ME_RE = re.compile(r"autiz|asperger|invalid|hendikep|\bgluv|nagluv|slijep|slabovid|nevidj|cerebraln|sindrom\w* daun|\bdaun|ometen|"
                   r"razvojn\w+ smetn|mentaln\w+ (smetn|zaostal)|epilep|distrofij|specijaln\w+ potreb|osob\w+ sa (invalid|posebnim)|paraplegi|\bparalizi", re.I)
CITY_MD = {"MUN.CHISINAU": "Chișinău", "MUN.BALTI": "Bălți"}


def latest(pkg, ext):
    d = rc.get_json(pkg)["result"]
    res = [r for r in d["resources"] if r["url"].lower().split("?")[0].endswith(ext)]
    res.sort(key=lambda r: r.get("created", ""))
    return res[-1]["url"]


def moldova():
    path = os.path.join(SCRATCH, "md.xlsx")
    if not os.path.exists(path):
        os.makedirs(SCRATCH, exist_ok=True)
        rc.download(latest(MD_PKG, ".xlsx"), path)
    geoc = rc.Geocoder(os.path.join(SCRATCH, "geocode_cache_md.jsonl"), countrycodes="md")
    out, nogeo = [], 0
    for r in openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(min_row=6, values_only=True):
        idno, _, form, name, _, addr, _, liquid = (list(r) + [None] * 8)[:8]
        if not name or liquid or not MD_RE.search(str(name)):
            continue
        if re.match(r"Sindicat|Organiza[tț]ia Sindical", str(name), re.I) or form and re.search(r"sindicat|religioas", str(form), re.I):
            continue
        name = rc.clean(name)
        m = re.match(r"(Asocia[tț]ie ob[sș]teasc[aă]|Institu[tț]ie public[aă]|Alte organiza[tț]ii necomerciale)\s+(?=[A-ZĂÂÎȘȚŞŢ\"“«])", name)
        if m:
            name = name[m.end():]
        name = re.sub(r"\s+", " ", name).strip()
        if name.startswith('"') and (name.endswith('"') or name.count('"') == 1):
            name = name.strip('"')
        name = name.strip(" ")
        a = rc.clean(addr)
        a = re.sub(r"^MD-\d+,\s*", "", a, flags=re.I)
        parts = [p.strip() for p in a.split(",") if p.strip()]
        parts = [p for p in parts if not re.match(r"(SEC\.|R-UL)", p, re.I)]
        flat = bool(re.search(r"ap\.\(of\.\)|\bap\.", a, re.I))
        loc = [p for p in parts if re.match(r"(MUN\.|OR\.|S\.|ORA[SȘ]|SAT)", p, re.I)] or parts[:1]
        town = CITY_MD.get(loc[0].upper(), re.sub(r"^(MUN\.|OR\.|S\.)", "", loc[0], flags=re.I).title()) if loc else ""
        street = next((p for p in parts if re.match(r"(str\.|bd\.|piața|piaţa|p-ța)", p, re.I)), "")
        number = ""
        if street:
            i = parts.index(street)
            if i + 1 < len(parts) and re.match(r"\d", parts[i + 1]):
                number = parts[i + 1]
        hit = None
        if street and not flat:
            hit = geoc.geocode(f"{street} {number}, {town}, Moldova".strip())
        if not hit:
            hit = geoc.geocode(f"{town}, Moldova")
        if not hit or not rc.plausible("MD", *hit):
            nogeo += 1
            continue
        address = ", ".join(x for x in ((f"{street} {number}".strip() if street and not flat else ""), town, "Moldova") if x)
        out.append(rc.entry(name, "Moldova Public Services Agency register of non-commercial organisations (dataset.gov.md)",
                            "advocacy", ["Advocacy", "Information & Support"], "Registered non-commercial organisation with a disability-related name.",
                            address=address, lat=hit[0], lng=hit[1], external_id=f"md-idno:{idno}",
                            scope=None if re.search(r"autis|asperger", name, re.I) else "general"))
    print(f"Moldova: {len(out)} entries, {nogeo} without a geocode")
    return out


def montenegro():
    path = os.path.join(SCRATCH, "me.xlsx")
    if not os.path.exists(path):
        os.makedirs(SCRATCH, exist_ok=True)
        rc.download(latest(ME_PKG, ".xlsx"), path)
    geoc = rc.Geocoder(os.path.join(SCRATCH, "geocode_cache_me.jsonl"), countrycodes="me")
    out, nogeo = [], 0
    for r in openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
        name, muni, addr, _, regno = (list(r) + [None] * 5)[:5]
        if not name or not ME_RE.search(str(name)) or ME_SKIP.search(str(name)):
            continue
        name, muni, addr = rc.clean(name), rc.clean(muni), rc.clean(addr)
        hit = geoc.geocode(f"{addr}, {muni}, Montenegro") if addr else None
        if not hit:
            hit = geoc.geocode(f"{muni}, Montenegro")
        if not hit or not rc.plausible("ME", *hit):
            nogeo += 1
            continue
        out.append(rc.entry(name.title() if name.isupper() else name, "Montenegro register of active NGOs (data.gov.me)", "advocacy",
                            ["Advocacy", "Information & Support"], "Registered non-governmental organisation with a disability-related name.",
                            address=", ".join(x for x in (addr.title() if addr else "", muni, "Montenegro") if x), lat=hit[0], lng=hit[1],
                            external_id=f"me-ngo:{regno}", scope=None if re.search(r"autiz|asperger", name, re.I) else "general"))
    print(f"Montenegro: {len(out)} entries, {nogeo} without a geocode")
    return out


def main():
    a = rc.finalize(os.path.join(rc.REPO, "new_resources_moldova.json"), moldova())
    b = rc.finalize(os.path.join(rc.REPO, "new_resources_montenegro.json"), montenegro())
    rc.spot_check(a, 15, fields=("name", "address"))
    rc.spot_check(b, 15, fields=("name", "address"))


if __name__ == "__main__":
    main()
